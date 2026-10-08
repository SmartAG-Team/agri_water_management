"""Check download identity, native-value preservation and monthly completeness."""
from pathlib import Path
import hashlib
import json
import netCDF4
import numpy as np
import pandas as pd
import xarray as xr
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]


def main():
    files=[json.loads(p.read_text()) for p in sorted((ROOT/'metadata/downloads').glob('*.json'))]
    assert len(files)==10
    assert not list((ROOT/'raw').rglob('*.part'))
    for r in files:
        path=ROOT/r['relative_path']
        assert path.stat().st_size==r['bytes']
        with path.open('rb') as f:assert hashlib.file_digest(f,'sha256').hexdigest()==r['sha256']
        if r['source_md5']:
            with path.open('rb') as f:assert hashlib.file_digest(f,'md5').hexdigest()==r['source_md5']
    profiles=[json.loads(p.read_text()) for p in sorted((ROOT/'metadata').glob('*_quality.json'))]
    assert len(profiles)==6
    product_checks=[]
    for r in profiles:
        name=r['product'];var=r['native_variable']
        audit=pd.read_csv(ROOT/'metadata'/f'{name}_time_axis.csv')
        assert audit.nominal_month.is_unique and audit.nominal_month.is_monotonic_increasing
        assert len(audit)==r['solution_count']
        with netCDF4.Dataset(ROOT/r['raw_path']) as raw, xr.open_dataset(
                ROOT/'derived/ncp_context'/f'{name}_NCP_context.nc',decode_times=False) as derived:
            assert len(raw['time'])==len(derived.time)==r['solution_count']
            assert np.array_equal(raw['time'][:],derived.time.values)
            assert list(derived.nominal_month.values)==list(audit.nominal_month)
            # Check all regional pixels for first, middle and last source solutions.
            lat_indices=np.array([np.flatnonzero(np.isclose(raw['lat'][:],v,rtol=0,atol=1e-8))[0] for v in derived.lat.values])
            lon_indices=np.array([np.flatnonzero(np.isclose(raw['lon'][:],v,rtol=0,atol=1e-8))[0] for v in derived.lon.values])
            it=[0,len(derived.time)//2,len(derived.time)-1]
            original=np.asarray(raw[var][it,lat_indices,lon_indices])
            output=derived[var].isel(time=it).transpose('time','lat','lon').values
            np.testing.assert_allclose(original,output,rtol=0,atol=0,equal_nan=True)
            assert float(derived.lat.min())>=28 and float(derived.lat.max())<=44
            assert float(derived.lon.min())>=109 and float(derived.lon.max())<=124
            assert derived[var].attrs['units'].startswith('cm')
        s=pd.read_csv(ROOT/'series'/f'{name}_regional_monthly.csv')
        assert len(s)==3*r['solution_count']
        assert not s.duplicated(['product','region','month']).any()
        assert np.isfinite(s.twsa_mm_as_released).all()
        assert s.valid_area_fraction.between(0,1).all()
        product_checks.append({'product':name,'months':r['solution_count'],
            'first':r['first_month'],'last':r['last_month'],
            'missing_within_span':len(r['missing_months_within_span']),
            'source_subset_values_identical':True})
    coverage=pd.read_csv(ROOT/'catalog/monthly_coverage_30years.csv')
    assert len(coverage)==360*6
    assert not coverage.duplicated(['product','month']).any()
    assert coverage.groupby('product').size().eq(360).all()
    monthly=pd.read_csv(ROOT/'series/monthly_1996-10_2026-09_by_product_region.csv')
    assert len(monthly)==360*6*3
    assert not monthly.duplicated(['product','region','month']).any()
    missing=~monthly.availability_status.isin(['satellite_solution','model_reconstruction'])
    assert monthly.loc[missing,'twsa_mm_as_released'].isna().all()
    assert monthly.loc[~missing,'twsa_mm_as_released'].notna().all()
    for name,expected in [('CSR_RL0603',259),('GSFC_RL06v20',255)]:
        c=coverage.loc[coverage['product']==name]
        assert (c.status=='satellite_solution').sum()==expected
        assert (c.status=='inter_mission_gap').sum()==11
        assert (c.status=='before_satellite_record').sum()==66
    g=next(r for r in profiles if r['product']=='GSFC_RL06v20')
    assert g['native_hdf_version']=='RL06 v2.0' and g['hdf_netcdf_time_axis_matched']
    with xr.open_dataset(ROOT/'derived/ncp_context/GSFC_RL06v20_native_mascons_with_uncertainty.nc') as d:
        assert len(d.time)==255 and len(d.mascon)==g['regional_native_mascons']
        assert all(v in d for v in ['noise_2sigma','leakage_2sigma','leakage_trend'])
    book=load_workbook(ROOT/'GRACE_monthly_NCP.xlsx',read_only=True,data_only=True)
    for name in ['NCPbox','six_provinces','NCP_2025_study']:
        sheet=book[name+'_TWSA_mm']
        assert sheet.max_row==361 and sheet.max_column==7
        assert sheet.cell(2,1).value=='1996-10' and sheet.cell(361,1).value=='2026-09'
    book.close()
    result={'status':'passed','source_files':len(files),'raw_bytes':sum(r['bytes'] for r in files),
        'publisher_md5_verified_files':sum(bool(r['source_md5']) for r in files),
        'products':product_checks,'requested_calendar_months':360,'coverage_rows':len(coverage),
        'regional_calendar_rows':len(monthly),'regions':3,'native_gsfc_mascons':g['regional_native_mascons'],
        'missing_values_not_interpolated':True,'raw_vs_subset_checks':'Every NCP-context grid cell in first, middle and last solutions; exact equality',
        'gsfc_native_uncertainty_components_preserved':True}
    (ROOT/'catalog/verification.json').write_text(json.dumps(result,indent=2))
    print(json.dumps(result,indent=2))


if __name__=='__main__':main()
