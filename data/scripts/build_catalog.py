"""Create the source catalogue and a local HTML index without changing raw data."""
from pathlib import Path
import json
import re
import html
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from prepare_collection import flattened

ROOT=Path(__file__).resolve().parents[2];DATA=ROOT/'data'
NOT_DATA={'metadata.json','files_metadata.json','source_terms.json','download_issue.txt'}
LABELS={'luancheng':'栾城','yucheng':'禹城','shangqiu':'商丘','fengqiu':'封丘','gucheng':'固城','cern_groundwater':'CERN多站地下水'}
EXTERNAL={
 'shandong_agricultural_university_irrigation_2012_2014':('山东农业大学','小麦补灌田间试验，2012–2014','土壤水分、根系、光合、器官干物质、灌浆过程','图中数值；不能视为完整逐日驱动与田间管理数据包','field_observations','downloaded'),
 'cau_wageningen_quzhou_survey':('中国农业大学／Wageningen','曲周；村级2017–2018，农户2020','农户和地块投入产出、管理、产量','调查数据；缺少连续土壤水分和ET过程观测','farm_survey','downloaded'),
 'cau_irrigation_districts':('中国农业大学相关研究；以仓库作者信息为准','中国371个灌区，2010–2017','灌溉量、灌溉面积、水利用效率和Budyko参数','ET为水量平衡推算；不作为独立实测ET真值','regional_statistics_and_derived','downloaded'),
 'peking_university_water_use':('北京大学相关团队','中国地市级，1965–2013；各变量年份不同','作物用水、灌溉面积、统计表、地市边界和代码','区域统计尺度；与田间灌溉量、抽水量和耗水量区分','regional_statistics','downloaded'),
 'luancheng_vadose_zone':('中科院栾城临界带研究团队','栾城48 m深竖井灌溉试验','深层包气带水分、水势与灌溉响应','用于深层渗漏/补给过程；与同站其他研究可能重叠','field_hydrology','downloaded'),
 'cas_eth_luancheng_recharge':('中科院／ETH合作研究团队','栾城地下水1974–2023；日水量表2021–2023','地下水埋深2274行；日降雨、灌溉、ET1095行；土壤与模拟结果','频率不均、井号待核实；观测和模拟结果共存；ET获取方法需核实；同站序列有重叠','mixed_observations_and_simulations','downloaded'),
 'cau_grass_maize_irrigation_2023_2024':('中国农业大学','华北平原，2023–2024','灌溉×氮肥、黑小麦/毛苕子–玉米、产量和土壤碳','前茬为黑小麦（triticale），不能当普通冬小麦；具体站点需由试验说明确认','field_treatment_observations','downloaded'),
 'henan_university_maize_spectra':('河南科技大学','玉米V6/V8/V12/R1，年份待核实','两品种、三氮水平的叶片光谱与SPAD','辅助氮/叶绿素模型验证；不是灌溉处理试验','field_leaf_observations','downloaded'),
 'ncp_long_term_nitrogen_yield':('仓库未明确机构；需核实','华北平原，目录2009–2025；已下载产量年份不连续','长期氮梯度田间产量及2025土壤碳','原表年份缺段；机构、站点及灌溉日程未核实；不计为独立灌溉试验','field_yield_and_soil','downloaded'),
 'groundwater_zenodo_17799537':('清华大学','华北平原，井点2005–2017；供用水2005–2023','891井点月度埋深、含水层类型、经纬度、供用水和代码','2018–2024原始井点需联系作者；部分观测可能与其他年鉴数据重复','well_observations_and_statistics','downloaded'),
 'groundwater_zenodo_7798617':('Water Resources Research 2025研究团队','华北平原；2005–2018/2005–2016','559行水位异常值序列、130行年鉴序列、井型和坐标','559和130是表行/序列数，并非观测次数；异常值基线和年鉴量定义需复核；与清华数据非完全独立','well_observations','downloaded'),
 'groundwater_zenodo_6378456':('首都师范大学','华北平原，2003–2016','452井地下水位趋势图','仓库仅提供PDF图，无逐井时间序列；不计入可分析数值数据包','figure_only','document_only'),
 'groundwater_zenodo_17616859':('中山大学／中南大学／GFZ／河海大学等','华北平原，地下水2018–2024','潜水/承压水月度埋深和形变','仓库明确限制访问；仅保存元数据，未下载数据','well_observations_restricted','restricted'),
 'groundwater_zenodo_15797080':('清华大学','Nature Communications 2025关联原版本','地下水恢复研究数据和代码','该版本限制访问；已取得公开v2（17799537），勿重复计数','older_restricted_version','restricted'),
 'yucheng_fluxnet_2011_2020':('ChinaFLUX／FLUXNET','禹城，2011–2020','半小时至年度碳、水、能量通量及质量标记','包含缺测填补和派生变量；使用QC筛选，与NESDC同站记录可能重叠','flux_observations_and_processed','downloaded'),
 'china_groundwater_grid_2005_2022':('辽宁师范大学／水科院／清华等；以论文署名为准','中国，2005–2022，月度1 km','地下水位IDW栅格；插值代码已下载','9.15 GB栅格未下载；属于空间插值产品，不是独立井点观测','interpolated_grid','code_only'),
}

def datafiles(folder):
    return [p for p in folder.iterdir() if p.is_file() and p.name not in NOT_DATA]

def build():
    rows=[]
    for p in sorted((DATA/'raw/nesdc').rglob('metadata.json')):
        d=flattened(json.loads(p.read_text())['data']);i=p.parent.name;site=p.parent.parent.name
        listing=json.loads((p.parent/'files_metadata.json').read_text())['data']
        fs=[p.parent/f['fileName'] for f in listing['data']]
        complete=listing['totalCount']==len(fs) and all(f.exists() for f in fs)
        term=json.loads((p.parent/'source_terms.json').read_text())
        lic='; '.join(l['name'] or l['url'] for l in term['licenses']) or '未提取到许可链接，参见源站'
        title=d['dataSetTitle'];notes='目录年份为数据集整体范围；具体变量/样地覆盖见工作表。独立验证需匹配样地、年份、品种及管理。'
        if site=='cern_groundwater':notes='FQA表存在异站编码和重复/冲突；LCA/YCA也有重复。见地下水质量表。YCA与新版禹城数据重叠。'
        if i=='67dd247d7e2817117c3b9efc':notes='地下水6272行、51条埋深空值、3条重复日期；土壤水分中子仪2005–2016、TDR2017–2022。'
        if i=='67dd19c67e2817117c3b9ef2':notes='已核实地面气象为逐日表；单位与辐射转换见说明文档。'
        if i=='67dd23167e2817117c3b9ef5':notes='产量表含多个样地和重复样方；按样地×年份聚合，不能把重复样方当独立年份。'
        if i in ['5fc9880d042ebb415eebe762','6020f966042ebb719ff186bf','5fc9868c042ebb415eebe73f','60e27c877e28173cf0d94a4c']:notes='历史整编气象表含月值；不是完整逐日作物模型驱动。另需匹配逐日气象。'
        text=title+' '+str(d.get('dataSetDesc',''))+' '+str(d.get('keyword',''))
        rows.append({'dataset_id':i,'title':title,'provider':'NESDC','institution':d.get('creator',''),'site':LABELS.get(site,site),
                     'period':d.get('temporalCoverage',''),'variables':d.get('keyword',''),'irrigation_mentioned':bool(re.search('灌溉|灌水',text)),
                     'groundwater_mentioned':bool(re.search('地下水',text)),'data_type':'station_observations','status':'downloaded' if complete else 'incomplete',
                     'original_data_files':len(fs),'original_bytes':sum(f.stat().st_size for f in fs if f.exists()),'doi':d.get('doi',''),
                     'source_url':'https://www.nesdc.org.cn/sdo/detail?id='+i,'license':lic,'folder':str(p.parent.relative_to(ROOT)),'limitations':notes})
    for p in sorted((DATA/'raw/external').glob('*/metadata.json')):
        key=p.parent.name;inst,period,variables,notes,kind,status=EXTERNAL[key];raw=json.loads(p.read_text())
        title=key;doi='';license='';url=''
        if 'snapshot' in raw:
            d=raw['snapshot'];title=d['name'];doi=d['doi'];license=d.get('licence',{}).get('short_name','');url=f'https://data.mendeley.com/datasets/{d["id"]}/{d["version"]}'
        elif 'metadata' in raw and 'doi' in raw:
            d=raw['metadata'];title=d['title'];doi=raw['doi'];license=d.get('license',{}).get('id','');url=raw['links']['self_html']
        elif 'files' in raw and 'title' in raw:
            title=raw['title'];doi=raw.get('doi','');license=raw.get('license',{}).get('name','');url=raw.get('url_public_html','')
        elif key=='yucheng_fluxnet_2011_2020':
            title='CN-Yuc FLUXNET 2011–2020 v1.3 r2';license='CC BY 4.0';url='https://meta.icos-cp.eu/objects/NFU16xgbyY-I0UsIppkbnnAV'
        fs=datafiles(p.parent)
        rows.append({'dataset_id':key,'title':title,'provider':'ICOS/FLUXNET' if key.startswith('yucheng_flux') else 'Zenodo' if 'zenodo' in key else 'Mendeley/Figshare',
                     'institution':inst,'site':'禹城' if key.startswith('yucheng_flux') else '栾城' if 'luancheng' in key else '区域/独立研究',
                     'period':period,'variables':variables,'irrigation_mentioned':'灌溉' in variables or '补灌' in period,'groundwater_mentioned':'groundwater' in key or '补给' in notes,
                     'data_type':kind,'status':status,'original_data_files':len(fs),'original_bytes':sum(f.stat().st_size for f in fs),
                     'doi':doi,'source_url':url,'license':license,'folder':str(p.parent.relative_to(ROOT)),'limitations':notes})
    frame=pd.DataFrame(rows);frame.to_csv(DATA/'catalog/datasets.csv',index=False,encoding='utf-8-sig')
    (DATA/'catalog/datasets.json').write_text(frame.to_json(orient='records',force_ascii=False,indent=2))
    tables={'数据集':frame,'工作表':pd.read_csv(DATA/'catalog/worksheets.csv'),'灌溉候选表':pd.read_csv(DATA/'catalog/irrigation_tables.csv'),
            '地下水质量':pd.read_csv(DATA/'catalog/groundwater_quality.csv'),'文件校验':pd.read_csv(DATA/'catalog/files.csv')}
    if (DATA/'curated/irrigation/station_irrigation_events_wheat_maize.csv').exists():
        tables['小麦玉米灌溉事件']=pd.read_csv(DATA/'curated/irrigation/station_irrigation_events_wheat_maize.csv')
    if (DATA/'catalog/literature_leads.csv').exists():tables['论文与数据线索']=pd.read_csv(DATA/'catalog/literature_leads.csv')
    with pd.ExcelWriter(DATA/'catalog/数据目录.xlsx',engine='openpyxl') as writer:
        for name,table in tables.items():
            table.to_excel(writer,sheet_name=name,index=False);ws=writer.sheets[name];ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='234E52')
            for column in ws.columns:
                ws.column_dimensions[column[0].column_letter].width=42 if column[0].value in ['title','source_file','limitations','variables','folder'] else 22
            for row in ws.iter_rows(min_row=2):
                for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
            for row in range(2,ws.max_row+1):ws.row_dimensions[row].height=42

    selected=['site','title','period','variables','status','limitations']
    headings=['站点/范围','数据集','时间范围','变量/用途','状态','使用限制']
    status={'downloaded':'已下载','restricted':'限制访问','document_only':'仅图件PDF','code_only':'仅代码，未下载栅格','incomplete':'未完整下载'}
    body=''
    for r in rows:
        body+='<tr>'+''.join('<td>'+('<a href="../'+html.escape(r['folder'])+'/">'+html.escape(str(r[c]))+'</a><br><a href="'+html.escape(r['source_url'])+'">源站</a>' if c=='title' else html.escape(status.get(str(r[c]),str(r[c]))))+'</td>' for c in selected)+'</tr>'
    page='''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>华北平原作物—灌溉—地下水数据目录</title><style>body{font:15px/1.6 system-ui,sans-serif;margin:32px;color:#193638}h1{font-size:25px}a{color:#16666e}input{font:inherit;padding:10px;width:min(700px,90%);margin:12px 0}table{border-collapse:collapse;width:100%}th,td{text-align:left;vertical-align:top;padding:12px;border-bottom:1px solid #dce7e6}th{background:#edf5f3;position:sticky;top:0}td:nth-child(2){min-width:210px}small{color:#556}tr[hidden]{display:none}</style><h1>华北平原作物—灌溉—地下水数据目录</h1><p>36个NESDC数据集、12个外部数值数据包已下载。另有1项仅提供图件、2项限制访问、1项仅下载插值代码。目录中的“灌溉/地下水提及”来自元数据，不代表每个变量均有连续观测。</p><p><a href="catalog/数据目录.xlsx">Excel总目录</a> · <a href="catalog/irrigation_tables.csv">灌溉候选工作表</a> · <a href="catalog/groundwater_quality.csv">地下水质量检查</a> · <a href="curated/groundwater/">地下水长表</a> · <a href="README.txt">数据说明</a></p><p>原始文件保持不变。CSV表格导出保留原表标题、单位和空行；地下水长表另附重复、缺测及站点编码异常标记。不同发布版本和同一井网的重用记录不能视为独立验证样本。</p><input aria-label="搜索数据集" placeholder="搜索：禹城、商丘、地下水、灌溉、清华……" oninput="document.querySelectorAll('tbody tr').forEach(r=>r.hidden=!r.textContent.toLowerCase().includes(this.value.toLowerCase()))"><table><thead><tr>'''+''.join('<th>'+h+'</th>' for h in headings)+'</tr></thead><tbody>'+body+'</tbody></table></html>'
    (DATA/'index.html').write_text(page)
    print(frame.groupby(['provider','status']).size().to_string())

if __name__=='__main__':build()
