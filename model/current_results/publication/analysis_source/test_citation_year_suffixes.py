"""Check that readers can distinguish cited records sharing author and year."""
import unittest

from docx import Document
from word_documents import Citations


class CitationYearSuffixTests(unittest.TestCase):
    def test_same_author_year_records_have_matching_unique_suffixes(self):
        groups = [
            'fengqiu2020_crop', 'fengqiu2020_soil_water',
            'yucheng2025_biology', 'yucheng2025_water',
            'chen2010_crop_productivity', 'chen2010_climate_water',
        ]
        blocks = [{'paragraphs': ['[CITE:' + key + ']' for key in groups]}]
        citations = Citations()
        citations.prepare(blocks)
        document = Document()
        for text in blocks[0]['paragraphs']:
            citations.add_text(document.add_paragraph(), text)
        self.assertEqual(
            [p.text for p in document.paragraphs],
            ['(Fengqiu Station, 2020a)', '(Fengqiu Station, 2020b)',
             '(Yucheng Station, 2025a)', '(Yucheng Station, 2025b)',
             '(Chen et al., 2010a)', '(Chen et al., 2010b)'])
        bibliography = citations.bibliography(document)
        for key, year in zip(groups, ['2020a', '2020b', '2025a', '2025b', '2010a', '2010b']):
            record = citations.records[key]
            row = next(row for row in bibliography if record['title'] in row)
            self.assertIn(year, row)
            self.assertIn(record['DOI'], row)


if __name__ == '__main__':
    unittest.main()
