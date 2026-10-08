"""Independently verify the chronological field partition, crop cards and full testing metrics."""
from pathlib import Path
from datetime import datetime,timezone
import argparse
import gzip
import hashlib
import json
import re
import sys
import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
sys.dont_write_bytecode=True
sys.path.insert(0,str(ROOT/'analysis_source'))
from audit_math import close,same_json,weights,moments
from field_source_identity import verify_field_targets


def independent_losses(station,field):
    rows=[]
    for (version,crop),g in station[station.split.eq('calibration')].groupby(['version','crop']):
        scales={}
        for variable,v in g.groupby('variable'):
            w=weights(v);mean=float(w@v.value.to_numpy(float))
            spread=float(np.sqrt(w@(v.value.to_numpy(float)-mean)**2))
            scales[variable]=max(.5 if variable in ['lai','et'] else 500.,spread)
        error=(g.predicted-g.value).to_numpy()/g.variable.map(scales).to_numpy()
        station_loss=.5*float(weights(g)@error**2)
        f=field[field.version.eq(version)&field.crop.eq(crop)&field.split.eq('calibration')]
        targets=[]
        for r in f.itertuples():
            for variable,o,p,scale in [
                ('seasonal_et',r.observed_et_mm,r.predicted_et_mm,max(50.,.15*r.observed_et_mm)),
                ('yield',r.yield_13pct_kg_ha*.87,r.grain_13pct_kg_ha*.87,max(1000.,.2*r.yield_13pct_kg_ha*.87))]:
                targets.append(dict(site=r.site,source_group_id=r.source_group_id,case_id=r.case_id,
                    variable=variable,error=(p-o)/scale))
        targets=pd.DataFrame(targets);field_loss=.35*float(weights(targets)@targets.error.to_numpy()**2)
        biomass=f.assign(variable='harvest_biomass')
        z=(biomass.predicted_biomass_kg_ha-biomass.observed_biomass_kg_ha).to_numpy()/np.maximum(1000.,.2*biomass.observed_biomass_kg_ha.to_numpy())
        biomass_loss=.15*float(weights(biomass)@z**2)
        p=g[g.variable.eq('et')].groupby(['site','source_group_id','case_id']).agg(
            observed=('value','sum'),simulated=('predicted','sum')).reset_index()
        p['variable']='observed_day_ET'
        z=(p.simulated-p.observed).to_numpy()/(.15*np.maximum(20.,p.observed.to_numpy()))
        matching_loss=.05*float(weights(p)@z**2)
        differences=[]
        for _,year in f.groupby('harvest_year'):
            lo,hi=year.set_index('treatment').loc[['W0','W3']].itertuples()
            differences.append((hi.predicted_et_mm-lo.predicted_et_mm)-(hi.observed_et_mm-lo.observed_et_mm))
        response_loss=.1*float(np.mean((np.array(differences)/50.)**2))
        rows.append(dict(version=version,crop=crop,station_loss=station_loss,field_loss=field_loss,
            matched_day_ET_loss=matching_loss,response_loss=response_loss,field_biomass_loss=biomass_loss,
            calibration_data_loss=station_loss+field_loss+matching_loss+response_loss+biomass_loss))
    return pd.DataFrame(rows)


def verify_primary_biomass():
    text=(ROOT/'source_snapshots/primary_methods_and_table1.txt').read_text()
    assert 'cut at the ground level' in text and 'oven-dried to a constant weight' in text
    start=text.index('L1085: Table 1');end=text.index('L1188: Fig. 3.',start)
    table={};year=None
    for line in text[start:end].splitlines():
        m=re.search(r'L\d+: (201[5-8])–(201[6-9])$',line.strip())
        if m:year=int(m.group(2));continue
        m=re.search(r'L\d+: (W[0-3]) (.*)',line)
        if not m:continue
        pairs=re.findall(r'(\d+(?:\.\d+)?)±(\d+(?:\.\d+)?)',m.group(2))
        assert len(pairs)==8
        for crop,index in [('wheat',3),('maize',7)]:
            table[crop,year,m.group(1)]=tuple(float(v)*1000. for v in pairs[index])
    data=pd.read_csv(ROOT/'data/confirmed_field_biomass.csv')
    assert len(data)==32 and data.case_id.is_unique and set(data.split)=={'calibration','validation'}
    for row in data.itertuples():
        mean,se=table[row.crop,int(row.harvest_year),row.treatment]
        close(row.observed_biomass_kg_ha,mean,'primary dry biomass mean')
        close(row.standard_error_kg_ha,se,'primary biomass standard error')
        assert row.field_replicates==3 and row.quantity=='aboveground_dry_biomass_including_grain'
    eligibility=json.loads((ROOT/'data/biomass_eligibility.json').read_text())
    assert eligibility['confirmed'] and not eligibility['biomass_moisture_conversion']
    return data


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--read-only',action='store_true');args=parser.parse_args()
    manifest=json.loads((ROOT/'verification/input_manifest.json').read_text())
    assert len({row['snapshot'] for row in manifest})==len(manifest)
    for row in manifest:
        assert hashlib.sha256((ROOT/row['snapshot']).read_bytes()).hexdigest()==row['sha256'],row['snapshot']
    applied=pd.read_csv(ROOT/'data/initialization_applied.csv').set_index('case_id')
    assert len(applied)==26 and applied.split.eq('calibration').sum()==11 and applied.split.eq('validation').sum()==15
    assert not applied.confirmed_sowing_day_state.any()
    for path in sorted((ROOT/'inputs/reference').rglob('*.json')):
        kind=path.parent.name
        case=path.name.removesuffix('_crop.json') if kind=='field' else path.stem
        current=json.loads(path.read_text())
        original=json.loads((ROOT/'source_snapshots/original_reference_inputs'/kind/path.name).read_text())
        if case in applied.index:
            row=applied.loc[case]
            close(current['inputs']['initial_theta'],json.loads(row.mapped_theta),'initialization rule applied')
            widths=np.array([l['thickness_mm'] for l in current['inputs']['soil_layers']])
            close(np.dot(current['inputs']['initial_theta'],widths),row.source_storage_mm,'measured storage conserved')
            month_end=(pd.Timestamp(int(row.profile_year),int(row.profile_month),1)+pd.offsets.MonthEnd(0)).date()
            assert month_end<pd.Timestamp(current['inputs']['start_date']).date()
            current['inputs']['initial_theta']=original['inputs']['initial_theta']
        assert same_json(current,original),'Only declared initial_theta changes are allowed'
    biomass=verify_primary_biomass().set_index('case_id')
    frozen=json.loads((ROOT/'parameters/frozen_model.json').read_text())
    assert not frozen['testing_used_for_selection'] and frozen['extra_free_parameters']==0
    assert frozen['individual_coefficient_priors_unchanged'] and frozen['testing_retrospective']
    for source,digest in frozen['fit_sha256'].items():assert hashlib.sha256((ROOT/source).read_bytes()).hexdigest()==digest
    for crop in ['wheat','maize']:
        card=frozen['candidates']['soil_refit'][crop]
        assert card['success'] and not card['testing_used'] and card['field_biomass_basis_confirmed']
        assert card['individual_coefficient_priors_unchanged']
        assert card['empirical_assimilation_water_response_retained']
        assert card['field_years']==[2016,2017,2018] and card['field_testing_years']==[2019]
        assert 'Wuqiao-2019' not in card['calibration_groups']
        assert card['free_parameters']==11 and card['extra_free_parameters']==0
        assert card['vector_order'][-1]=='root_compensation_fraction'
        assert card['lower_bounds'][-1]==0. and card['upper_bounds'][-1]==1.
        assert card['upper_bounds'][0]==1.5
        assert card['upper_bounds'][4]==(1.4 if crop=='wheat' else 1.5)
        if crop=='maize':assert card['upper_bounds'][5]==1.5
        assert np.all(np.array(card['vector'])>=card['lower_bounds']) and np.all(np.array(card['vector'])<=card['upper_bounds'])
        used=pd.read_csv(ROOT/'data'/('used_calibration_field_'+crop+'.csv'))
        assert len(used)==36 and used.case_id.nunique()==12 and used.split.eq('calibration').all()
        assert set(used.harvest_year)=={2016.,2017.,2018.}
        b=used[used.variable.eq('harvest_biomass')]
        close(b.value,biomass.loc[b.case_id,'observed_biomass_kg_ha'],'fitted dry biomass without moisture conversion')
    observations=pd.read_csv(ROOT/'data/observations.csv',low_memory=False)
    inventory=pd.read_csv(ROOT/'data/case_inventory.csv')
    assert len(observations)==observations.observation_id.nunique()==7397 and len(inventory)==132
    assert len(observations[observations.split.eq('calibration')])==3771
    assert inventory.groupby('source_group_id').split.nunique().eq(1).all()
    station=pd.read_csv(ROOT/'predictions/station_comparisons.csv',low_memory=False)
    field=pd.read_csv(ROOT/'predictions/field_comparisons.csv')
    summary=pd.read_csv(ROOT/'predictions/case_summaries.csv')
    original_field=pd.read_csv(ROOT/'source_snapshots/R59_field_comparisons.csv')
    original_field=original_field[original_field.version.eq('compensated_refit')]
    verify_field_targets(field,original_field)
    verify_field_targets(pd.read_csv(ROOT/'data/reference_field_comparisons.csv'),original_field)
    assert field.loc[field.harvest_year.eq(2019),'split'].eq('validation').all()
    assert field.loc[field.harvest_year.isin([2016,2017,2018]),'split'].eq('calibration').all()
    assert field.groupby(['version','source_group_id']).split.nunique().eq(1).all()
    assert len(field)==64 and len(summary)==328 and set(field.version)=={'reference','soil_refit'}
    oldq=pd.read_csv(ROOT/'data/reference_station_comparisons.csv',low_memory=False)
    oldq=oldq[oldq.version.eq('supplied_fit')].set_index('observation_id')
    oldf=pd.read_csv(ROOT/'data/reference_field_comparisons.csv')
    oldf=oldf[oldf.version.eq('supplied_fit')].set_index('case_id')
    for version,g in station.groupby('version'):
        assert len(g)==g.observation_id.nunique()==7397 and g.site.nunique()==5
        assert g.groupby('source_group_id').split.nunique().eq(1).all()
        for column in ['case_id','site','crop','split','variable','value','window_start','window_end']:
            pd.testing.assert_series_equal(g[column].reset_index(drop=True),oldq.loc[g.observation_id,column].reset_index(drop=True),check_names=False,check_dtype=False)
        if version=='reference':
            expected=oldq.loc[g.observation_id,'predicted'].to_numpy()
            eligible=np.isfinite(expected)
            close(g.predicted.to_numpy()[eligible],expected[eligible],'station reference replay under same initialization')
    for column in ['predicted_et_mm','predicted_biomass_kg_ha','grain_13pct_kg_ha']:
        g=field[field.version.eq('reference')]
        close(g[column],oldf.loc[g.case_id,column],'field reference replay '+column)
    for row in field.itertuples():close(row.observed_biomass_kg_ha,biomass.loc[row.case_id,'observed_biomass_kg_ha'],'matched field observation identity')
    trajectories=windows=0
    for path in sorted((ROOT/'inputs/resolved').rglob('*.json')):
        kind,version=path.parent.name,path.parent.parent.name
        case=path.name.removesuffix('_crop.json') if kind=='field' else path.stem
        p=json.loads(path.read_text());original=json.loads((ROOT/'inputs/reference'/kind/path.name).read_text())
        assert same_json(p['inputs'],original['inputs']),('prepared soil initialization, weather, calendar and management',case,version)
        if p.get('presowing'):assert same_json(p['presowing']['inputs'],original['presowing']['inputs'])
        if version=='reference':assert same_json(p,original)
        else:
            raw=json.loads((ROOT/'inputs'/kind/path.name).read_text());crop=p['inputs']['crop']
            card=frozen['candidates']['soil_refit'][crop];x=card['vector']
            c,b,h=p['parameters']['crop'],raw['parameters']['crop'],p['parameters']['hydrology']
            close(c['transpiration_coefficient'],min(1.5,b['transpiration_coefficient']*x[0]),'transpiration fit')
            close(c['soil_evaporation_coefficient'],x[1],'soil evaporation fit')
            close(h['readily_available_water_fraction'],x[2],'root-zone depletion fit')
            close(c['profile']['rue_g_mj'],b['profile']['rue_g_mj']*x[3],'RUE fit')
            if crop=='wheat':
                grain,stress,decay,depth=5,6,7,8
                for key in ['leaf_allocation_multiplier','early_leaf_allocation_multiplier','late_leaf_allocation_multiplier']:
                    if key in b['growth_process']:close(c['growth_process'][key],min(2.,b['growth_process'][key]*x[4]),'wheat leaf allocation')
                close(h['root_extraction_fraction_day'],10**x[9],'wheat extraction rate')
            else:
                grain,stress,decay,depth=6,7,8,9
                for key,index in [('early_leaf_allocation_multiplier',4),('late_leaf_allocation_multiplier',5)]:
                    close(c['growth_process'][key],min(2.,b['growth_process'][key]*x[index]),'maize leaf allocation')
                close(h['root_extraction_fraction_day'],1.,'maize extraction rate')
            close(c['profile']['kernels_per_g_flowering_biomass'],b['profile']['kernels_per_g_flowering_biomass']*x[grain],'grain number')
            close(c['growth_process']['assimilation_water_stress_exponent'],x[stress],'empirical assimilation-water response')
            assert 'transpiration_efficiency_kpa' not in c['growth_process']
            close(h['root_density_decay_m_inv'],x[decay],'root density')
            close(c['profile']['root_max_mm'],min(x[depth],sum(l['thickness_mm'] for l in p['inputs']['soil_layers'])),'root depth')
            assert h['matric_interface_method']=='relative_arithmetic' and h['plant_water_stress_method']=='compensated_layer_depletion'
            close(h['root_compensation_fraction'],x[10],'bounded root compensation')
            assert len(x)==11 and 0.<=x[10]<=1.
            if p.get('presowing'):assert same_json(p['presowing']['parameters'],p['parameters'])
        with gzip.open(ROOT/'predictions/full_daily'/version/kind/(case+'.json.gz'),'rt') as stream:d=pd.DataFrame(json.load(stream))
        assert d.date.tolist()==[row['date'] for row in p['inputs']['weather']]
        assert d.balance_residual_mm.abs().max()<1e-6 and d.crop_carbon_residual_kg_ha.abs().max()<1e-6
        if version=='soil_refit':
            assert np.all(d.growth_kg_ha<=d.potential_growth_kg_ha*d.nutrition_factor+1e-8)
        close(d.et_mm,d.transpiration_mm+d.soil_evaporation_mm+d.canopy_evaporation_mm,'actual ET component closure')
        close(d.storage_final_mm-d.storage_initial_mm,d.precipitation_mm+d.irrigation_field_mm+d.capillary_rise_mm-d.et_mm-d.runoff_mm-d.bottom_drainage_mm,'daily water budget')
        theta=np.asarray(d.soil_theta.tolist());layers=p['inputs']['soil_layers']
        assert np.all(theta>=np.array([l['air_dry'] for l in layers])-1e-9)
        assert np.all(theta<=np.array([l['saturation'] for l in layers])+1e-9)
        row=summary[summary.version.eq(version)&summary.case_id.eq(case)].iloc[0]
        close(d.et_mm.sum(),row.predicted_et_mm,'seasonal actual ET')
        close(d.yield_kg_ha.iloc[-1]/.87,row.grain_13pct_kg_ha,'13-percent grain moisture')
        close(d.biomass_kg_ha.iloc[-1],row.predicted_biomass_kg_ha,'harvest dry biomass')
        if kind=='station':
            g=station[station.version.eq(version)&station.case_id.eq(case)]
            for obs in g.itertuples():
                column={'lai':'lai','biomass':'biomass_kg_ha','harvest_biomass':'biomass_kg_ha','yield':'yield_kg_ha','et':'et_mm'}[obs.variable]
                selected=d[d.date.ge(obs.window_start)&d.date.le(obs.window_end)]
                assert len(selected)==(pd.Timestamp(obs.window_end)-pd.Timestamp(obs.window_start)).days+1
                close(selected[column].mean(),obs.predicted,'matched observation window');windows+=1
        trajectories+=1
    assert trajectories==328 and windows==14794
    metric_rows=0
    for name,by_site in [('station_metrics.csv',False),('site_metrics.csv',True)]:
        for row in pd.read_csv(ROOT/'tables'/name).itertuples():
            g=station[station.version.eq(row.version)&station.crop.eq(row.crop)&station.split.eq(row.split)&station.variable.eq(row.variable)]
            if by_site:g=g[g.site.eq(row.site)]
            expected=moments(g.value,g.predicted,weights(g))
            for column in ['rmse','bias','nse','r_squared']:close(getattr(row,column),expected[column],'balanced station metric')
            metric_rows+=1
    contrasts=pd.read_csv(ROOT/'tables/field_ET_contrasts.csv')
    for row in contrasts.itertuples():
        g=field[field.version.eq(row.version)&field.crop.eq(row.crop)&field.split.eq(row.split)&field.harvest_year.eq(row.harvest_year)].set_index('treatment')
        close(row.observed_mm,g.loc['W3','observed_et_mm']-g.loc['W0','observed_et_mm'],'observed ET contrast')
        close(row.predicted_mm,g.loc['W3','predicted_et_mm']-g.loc['W0','predicted_et_mm'],'predicted ET contrast')
    for row in pd.read_csv(ROOT/'tables/field_metrics.csv').itertuples():
        if row.variable=='seasonal_et_response':
            g=contrasts[contrasts.version.eq(row.version)&contrasts.crop.eq(row.crop)&contrasts.split.eq(row.split)]
            expected=moments(g.observed_mm,g.predicted_mm);expected.update(nrmse_percent=np.nan)
            if len(g)<3:expected['r_squared']=np.nan
        else:
            g=field[field.version.eq(row.version)&field.crop.eq(row.crop)&field.split.eq(row.split)]
            a,b={'seasonal_et':('observed_et_mm','predicted_et_mm'),'yield':('yield_13pct_kg_ha','grain_13pct_kg_ha'),
                 'biomass':('observed_biomass_kg_ha','predicted_biomass_kg_ha')}[row.variable]
            expected=moments(g[a],g[b])
        for column in ['rmse','bias','nrmse_percent','nse','r_squared']:close(getattr(row,column),expected[column],'field metric')
        metric_rows+=1
    independent=independent_losses(station,field)
    loss=pd.read_csv(ROOT/'tables/calibration_data_losses.csv')
    for row in loss.itertuples():
        expected=independent[independent.version.eq(row.version)&independent.crop.eq(row.crop)].iloc[0]
        for column in ['station_loss','field_loss','matched_day_ET_loss','response_loss','field_biomass_loss','calibration_data_loss']:
            close(getattr(row,column),expected[column],'independent calibration loss '+column)
    means=independent.groupby('version').calibration_data_loss.mean()
    assert frozen['selected_version']==means.idxmin()
    for version,value in means.items():close(value,frozen['mean_crop_calibration_data_loss'][version],'frozen calibration selection')
    checks=pd.read_csv(ROOT/'tables/seasonal_ET_working_checks.csv')
    metrics=pd.read_csv(ROOT/'tables/field_metrics.csv')
    for row in checks.itertuples():
        g=metrics[metrics.version.eq(row.version)&metrics.crop.eq(row.crop)&metrics.split.eq(row.split)].set_index('variable')
        f=field[field.version.eq(row.version)&field.crop.eq(row.crop)&field.split.eq(row.split)]
        e=g.loc['seasonal_et'];expected=dict(RMSE_le50=e.rmse<=50,nRMSE_le15=e.nrmse_percent<=15,
            absolute_bias_le10pct=abs(e.bias)<=.1*f.observed_et_mm.mean(),NSE_positive=e.nse>0,
            response_RMSE_le50=g.loc['seasonal_et_response','rmse']<=50)
        for k,v in expected.items():assert getattr(row,k)==v
        assert row.working_ET_checks_pass==all(expected.values())
        assert row.confirmed_field_yield_nRMSE_le20==(g.loc['yield','nrmse_percent']<=20)
        assert not row.field_biomass_is_conditional
    water=pd.read_csv(ROOT/'tables/field_soil_water_diagnostics.csv')
    assert len(water)==64 and water.diagnostic_not_independent_validation.all()
    assert not water.source_sampling_dates_confirmed.any()
    assert not water.independent_of_source_water_balance_ET.any()
    for row in water.itertuples():
        with gzip.open(ROOT/'predictions/full_daily'/row.version/'field'/(row.case_id+'.json.gz'),'rt') as stream:d=pd.DataFrame(json.load(stream))
        close(row.simulated_initial_storage_mm,d.soil_storage_initial_mm.iloc[0],'soil storage initialization diagnostic')
        close(row.simulated_final_storage_mm,d.soil_storage_final_mm.iloc[-1],'soil storage endpoint diagnostic')
        close(row.simulated_depletion_mm,d.soil_storage_initial_mm.iloc[0]-d.soil_storage_final_mm.iloc[-1],'soil storage depletion diagnostic')
    book=load_workbook(ROOT/'tables/validation_data_and_results.xlsx',read_only=True,data_only=True)
    assert {'Source_SHA256','Confirmed_field_biomass','Calibration_losses','Working_accuracy_checks','Exact_used_weather'}<=set(book.sheetnames)
    sha=list(book['Source_SHA256'].values)
    assert sha[1:]==[(r['source'],r['snapshot'],r['sha256']) for r in manifest]
    assert book['Confirmed_field_biomass'].max_row==33 and book['Field_predictions'].max_row==65
    book.close()
    if (ROOT/'verification/artifact_manifest.json').exists():
        for row in json.loads((ROOT/'verification/artifact_manifest.json').read_text()):
            assert hashlib.sha256((ROOT/row['path']).read_bytes()).hexdigest()==row['sha256'],row['path']
    chosen=checks[checks.version.eq(frozen['selected_version'])]
    achieved=bool(chosen.working_ET_checks_pass.all() and chosen.confirmed_field_yield_nRMSE_le20.all())
    receipt=dict(input_snapshots=len(manifest),confirmed_field_biomass_records=32,trajectories=trajectories,
        observation_windows=windows,independent_metric_rows=metric_rows,reference_replayed_for_available_same_initialization_expectations=True,
        native_water_and_carbon_budgets_closed=True,whole_site_years_preserved=True,all_eligible_sites_retained=True,
        original_weather_soil_properties_and_irrigation_preserved=True, initialization_rule_applied_identically_to_reference_and_refit=True,calibration_selection_precedes_testing=True,
        testing_is_retrospective=True,scientific_goal_achieved=achieved,all_checks_passed=True,
        verified_at_utc=datetime.now(timezone.utc).isoformat())
    if not args.read_only:(ROOT/'verification/independent_checks.json').write_text(json.dumps(receipt,indent=2)+'\n')
    print(json.dumps(receipt,indent=2))


if __name__=='__main__':main()
