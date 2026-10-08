"""Check scientific sources, current paper bindings and deliverable integrity."""
from pathlib import Path
from datetime import datetime, timezone
import hashlib
import json
import re
import zipfile

from lxml import etree
import numpy as np
import pandas as pd
import pymupdf as fitz

ROOT = Path(__file__).resolve().parents[1]
PUB = ROOT / 'publication'
W = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def sha(path):
    h = hashlib.sha256()
    with path.open('rb') as f:
        for part in iter(lambda: f.read(4 * 1024 * 1024), b''):
            h.update(part)
    return h.hexdigest()


def main():
    checks = []
    def check(name, value):
        checks.append({'check': name, 'passed': bool(value)})
        if not value:
            raise ValueError(name)

    def load(path):
        return json.loads(path.read_text())

    copied = load(ROOT / 'verification/source_consolidation.json')
    for name, digest in copied['copied_files_sha256'].items():
        check('Preserved calibration artifact ' + name, sha(ROOT / 'calibration' / name) == digest)
    manifest = load(ROOT / 'regional/verification/file_manifest.json')
    for name, record in manifest.items():
        file = ROOT / 'regional' / name
        check('Sealed regional artifact ' + name, file.is_file() and sha(file) == record['sha256'])
    native = load(ROOT / 'regional/verification/native_source_identity.json')
    for name, digest in native['files'].items():
        for run in ['calibration', 'regional']:
            folder = 'source_snapshots/native' if run == 'calibration' else 'source_snapshots/native_process'
            check(f'Identical {run} native engine ' + name, sha(ROOT / run / folder / name) == digest)
    for run in ['calibration', 'regional']:
        inputs = load(ROOT / run / 'verification/input_manifest.json')
        for record in inputs:
            file = ROOT / run / record['snapshot']
            check(f'Exact {run} input ' + record['snapshot'], sha(file) == record['sha256'])

    frozen = load(ROOT / 'calibration/parameters/frozen_model.json')
    check('Selected current model', frozen['selected_version'] == 'management_refit')
    check('Selection excludes retrospective testing', not frozen['testing_used_for_selection'])
    check('2016–2018 water calibration; 2019 retrospective testing',
          frozen['field_calibration_years'] == [2016, 2017, 2018] and frozen['field_testing_years'] == [2019])
    binding = load(PUB / 'verification/current_publication_binding.json')
    for name, digest in binding['source_sha256'].items():
        check('Exact publication source ' + name, sha(ROOT / name) == digest)
    transfer = load(ROOT / 'regional/verification/parameter_transfer.json')
    card = pd.read_csv(PUB / 'tables/shared_crop_parameters.csv').set_index('Parameter')
    for crop in ['wheat', 'maize']:
        rue = transfer['effective_parameters'][crop]['crop']['profile']['rue_g_mj']
        check('Current RUE transferred to Table S3 ' + crop, np.isclose(card.loc['Radiation-use efficiency', crop.title()], rue))
    check('No maximum LAI calibration parameter', not any('Maximum leaf area index' in s for s in card.index))
    check('Stage-specific SLA and senescence', sum('Specific leaf area at DVS' in s for s in card.index) == 3 and
          sum('Leaf senescence rate' in s for s in card.index) == 3)

    yearly = pd.read_csv(ROOT / 'regional/tables/regional_policy_annual_results.csv')
    means = yearly[yearly.harvest_year.between(2014, 2025)].groupby('policy').mean(numeric_only=True)
    gain = (means.loc['targeted_50pct', 'grain_production_t'] - means.loc['uniform_50pct', 'grain_production_t']) / 1e6
    et = (means.loc['targeted_50pct', 'modeled_total_et_volume_m3'] - means.loc['uniform_50pct', 'modeled_total_et_volume_m3']) / 1e9
    check('Current regional dry-grain arithmetic', np.isclose(gain, binding['grain_advantage_Mt']))
    check('Current regional ET arithmetic', np.isclose(et, binding['ET_advantage_km3']))
    controls = load(ROOT / 'regional/verification/matched_information_controls.json')
    check('Matched information controls have identical decisions and outcomes',
          all(v['decisions_identical'] and max(v['maximum_output_differences'].values()) == 0 for v in controls.values()))
    results = load(ROOT / 'regional/verification/current_publication_checks.json')
    check('Numerical water, carbon, continuity and allocation checks pass', results['all_checks_passed'])

    documents = {name: load(PUB / 'analysis_source' / (name + '_blocks.json'))
                 for name in ['manuscript', 'supplementary']}
    all_text = '\n'.join(str(b) for d in documents.values() for b in d['blocks'])
    result_blocks=[]
    active=False
    for block in documents['manuscript']['blocks']:
        heading=block.get('heading','')
        if heading.startswith('3.'):active=True
        elif heading.startswith('4.'):active=False
        if active:result_blocks.append(block)
    check('Results prose, captions and table notes contain no literature citations',
          result_blocks and all('[CITE:' not in str(block) for block in result_blocks))
    check('No obsolete independent Wuqiao claim', not re.search(r'independent Wuqiao|independent irrigation experiment', all_text, re.I))
    for old in ['0.358 Mt', '0.115 km', '97.26%', '98.66%', '225.72 mm', '12.40 mm', '2.729', '3.066', '0.97156']:
        check('No stale published value ' + old, old not in all_text)
    check('No obsolete supplementary references', not re.search(r'\bTable S(?:10|11|12)\b|\bFigure S12\b', all_text))
    check('Common crop model name', 'crop–soil model' not in all_text and 'crop-soil model' not in all_text)
    used_refs = set(re.findall(r'\[CITE:([^\]]+)\]', all_text))
    used_refs = {key for group in used_refs for key in group.split('|')}
    exported = load(PUB / 'literature/manuscript_references.csl.json')
    check('Exact cited-reference export', used_refs == {r['id'] for r in exported} and len(exported) == 41)
    abstract = next(b for b in documents['manuscript']['blocks'] if b.get('heading') == 'Abstract')['paragraphs'][0]
    check('Abstract word limit', len(abstract.split()) == binding['abstract_words'] <= 250)
    abstract_structure = load(PUB / 'verification/abstract_structure_20261008.json')
    components = abstract_structure['components']
    check('Abstract follows importance, gap, overview, results and implications',
          [c['role'] for c in components] == ['importance', 'gap', 'overview', 'results', 'implications'] and
          ' '.join(c['text'] for c in components) == abstract)
    check('Abstract has distinct importance and research-gap opening sentences',
          all(c['text'].endswith('.') and len(re.findall(r'\.(?:\s|$)', c['text'])) == 1
              for c in components[:2]))
    check('Abstract uses current grain and ET results', f'{gain:.2f}' in abstract and f'{et:.2f}' in abstract)
    check('Grammar-only AI declaration', next(b for b in documents['manuscript']['blocks'] if
          b.get('heading', '').startswith('Declaration of generative AI'))['paragraphs'] == ['AI assisted with grammar checking.'])
    visual=load(PUB/'verification/visual_revision_20261008.json')
    for figure in visual['figures']:
        file=PUB/figure['figure']
        check('Reviewed figure layout PNG '+figure['figure'],sha(file)==figure['png_sha256'])
        check('Reviewed figure layout PDF '+figure['figure'],sha(file.with_suffix('.pdf'))==figure['pdf_sha256'])
        for source,digest in figure['source_sha256'].items():
            check('Unchanged reviewed figure data '+figure['figure']+'/'+source,sha(ROOT/source)==digest)
    panels=load(PUB/'verification/panel_title_removal_20261008.json')
    check('All publication figures covered by panel-title removal',
          len(panels['figures'])==19 and panels['main_figures']==8 and panels['supplementary_figures']==11 and
          not panels['descriptive_panel_titles'])
    for figure in panels['figures']:
        file=PUB/figure['figure']
        check('Panel-letter layout PNG '+figure['figure'],sha(file)==figure['png_sha256'])
        check('Panel-letter layout PDF '+figure['figure'],sha(file.with_suffix('.pdf'))==figure['pdf_sha256'])
        with fitz.open(file.with_suffix('.pdf')) as vector:
            vector_text='\n'.join(page.get_text() for page in vector)
        check('Descriptive panel titles removed '+figure['figure'],
              not figure['descriptive_panel_titles'] and
              all(title not in vector_text for title in figure.get('removed_panel_titles',[])) and
              all(label in vector_text for label in figure.get('panel_labels',[])))
        for source,digest in figure['source_sha256'].items():
            check('Exact panel-letter layout source '+figure['figure']+'/'+source,sha(ROOT/source)==digest)
    cover = load(PUB / 'analysis_source/cover_letter_paragraphs.json')
    from docx import Document
    check('Concise cover letter matches its source and current word count',
          0 < len(' '.join(cover).split()) <= 300 and
          len(' '.join(cover).split()) == binding['cover_words'] and
          [p.text for p in Document(PUB/'documents/cover_letter.docx').paragraphs] == cover)
    audit=load(ROOT/'verification/2026-10-08_final_results_numbers_audit.json')
    check('Every final Results/Conclusions/Highlights number audited',audit['all_quantitative_claims_passed'] and
          audit['claim_count']==405 and audit['failed_claim_count']==0 and not audit['unmapped_claims'] and
          audit['manuscript_sha256']==sha(PUB/'analysis_source/manuscript_blocks.json'))
    highlights=load(PUB/'analysis_source/highlights_paragraphs.json')
    check('Five numbered-result highlights within publication length',len(highlights)==5 and
          all(len(x)<=85 and re.search(r'\d',x) for x in highlights))
    from docx import Document
    check('Standalone highlights match reviewed source', [q.text for q in Document(PUB/'documents/highlights.docx').paragraphs if q.text.strip()]==highlights)
    expected = {'manuscript': (8, 5), 'supplementary': (11, 9)}
    media = []
    paper_book=pd.ExcelFile(PUB/'tables/paper_data_and_tables.xlsx')
    for name, d in documents.items():
        figures = [b for b in d['blocks'] if 'figure' in b]
        tables = [b for b in d['blocks'] if 'table' in b]
        check(name + ' final figure/table counts', (len(figures), len(tables)) == expected[name])
        prefix = '' if name == 'manuscript' else 'S'
        check(name + ' figure labels sequential', [re.match(r'Figure (S?\d+)\.', b['caption']).group(1) for b in figures] ==
              [prefix + str(i) for i in range(1, len(figures) + 1)])
        check(name + ' table labels sequential', [re.match(r'Table (S?\d+)\.', b['table_caption']).group(1) for b in tables] ==
              [prefix + str(i) for i in range(1, len(tables) + 1)])
        for block in figures:
            path = PUB / block['figure']
            check('Current PNG/PDF pair ' + block['figure'], path.is_file() and path.with_suffix('.pdf').is_file())
            media.append(sha(path))
        for block in tables:
            check('Current table source ' + block['table'], (PUB / block['table']).is_file())
            label=re.match(r'Table (S?\d+)\.',block['table_caption']).group(1)
            source=pd.read_csv(PUB/block['table'])
            sheet=(label+'_'+Path(block['table']).stem)[:31]
            workbook=pd.read_excel(paper_book,sheet_name=sheet)
            check('Exact workbook paper table '+label,list(source.columns)==list(workbook.columns) and len(source)==len(workbook))
            for column in source:
                if pd.api.types.is_numeric_dtype(source[column]):
                    check('Workbook numeric column '+label+'/'+column,np.allclose(source[column].to_numpy(float),workbook[column].to_numpy(float),equal_nan=True))
                else:check('Workbook labels '+label+'/'+column,source[column].fillna('').astype(str).tolist()==workbook[column].fillna('').astype(str).tolist())

    doc_receipts = {}
    for name, count, tables in [('manuscript', 8, 5), ('supplementary_material', 11, 9), ('manuscript_package', 19, 14)]:
        path = PUB / 'documents' / (name + '.docx')
        with zipfile.ZipFile(path) as z:
            xml = etree.fromstring(z.read('word/document.xml'))
            math_ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
            math_count=len(xml.xpath('//m:oMath',namespaces=math_ns))
            check(name+' editable inline equations',math_count=={'manuscript':8,'supplementary_material':40,'manuscript_package':48}[name])
            if name in ['manuscript','manuscript_package']:
                inside=False;result_fields=[]
                for element in xml.find('w:body',W):
                    text=''.join(element.xpath('.//w:t/text()',namespaces=W))
                    if text=='3. Results':inside=True
                    elif text=='4. Discussion':inside=False
                    if inside:result_fields.extend(element.xpath('.//w:instrText/text()',namespaces=W))
                check(name+' Results contain no native citation fields',not any('CSL_CITATION' in s for s in result_fields))
            instructions = xml.xpath('//w:instrText/text()', namespaces=W)
            citations = [json.loads(s.split('CSL_CITATION ', 1)[1].strip()) for s in instructions if 'CSL_CITATION ' in s]
            ids = [c['citationID'] for c in citations]
            check(name + ' unique Zotero citation IDs', len(ids) == len(set(ids)))
            check(name + ' citation metadata retained', all('itemData' in i for c in citations for i in c['citationItems']))
            check(name + ' Zotero bibliography field retained', any('ZOTERO_BIBL' in s for s in instructions))
            check(name + ' balanced native Word fields', len(xml.xpath('//w:fldChar[@w:fldCharType="begin"]', namespaces=W)) ==
                  len(xml.xpath('//w:fldChar[@w:fldCharType="separate"]', namespaces=W)) == len(xml.xpath('//w:fldChar[@w:fldCharType="end"]', namespaces=W)))
            actual_media = [hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')]
            needed = media[:8] if name == 'manuscript' else media[8:] if name == 'supplementary_material' else media
            check(name + ' current embedded figure bytes', len(actual_media) == count and set(actual_media) == set(needed))
            tab = xml.xpath('//w:tbl', namespaces=W)
            check(name + ' all tables present', len(tab) == tables)
            for index, t in enumerate(tab):
                b = t.find('w:tblPr/w:tblBorders', W)
                check(f'{name} three-rule table {index + 1}', b is not None and
                      all(b.find('w:' + k, W).get('{'+W['w']+'}val') == 'single' for k in ['top', 'bottom']) and
                      all(b.find('w:' + k, W).get('{'+W['w']+'}val') == 'nil' for k in ['left', 'right', 'insideH', 'insideV']) and
                      bool(t.xpath('./w:tr[1]/w:tc/w:tcPr/w:tcBorders/w:bottom[@w:val="single"]', namespaces=W)))
            if name != 'manuscript':
                text = ''.join(xml.xpath('//w:t/text()', namespaces=W))
                check(name + ' tiny parameter scientific notation', '3.41 × 10⁻¹⁰' in text and '2.27 × 10⁻¹³' in text)
        with fitz.open(PUB / 'documents' / (name + '.pdf')) as pdf:
            text = '\n'.join(p.get_text() for p in pdf)
            check(name + ' reading PDF content', len(text) > 10000 and 'References' in text)
            for figure in range(1, count + 1):
                label = ('Figure ' + str(figure) + '.') if name == 'manuscript' else (
                    'Figure S' + str(figure) + '.' if name == 'supplementary_material' else
                    'Figure ' + str(figure) + '.' if figure <= 8 else 'Figure S' + str(figure - 8) + '.')
                check(name + ' PDF caption ' + label, label in text)
            pages = len(pdf)
        doc_receipts[name] = {'docx_sha256': sha(path), 'pdf_sha256': sha(path.with_suffix('.pdf')), 'pdf_pages': pages,
                              'citation_fields': len(citations), 'unique_citation_ids': True, 'embedded_current_figures': count}
    with fitz.open(PUB / 'documents/cover_letter.pdf') as pdf:
        check('Cover reading PDF one page', len(pdf) == 1)
    for name, n in [('main_figures', 8), ('supplementary_figures', 11)]:
        with fitz.open(PUB / 'documents' / (name + '.pdf')) as pdf:
            check(name + ' current figure collection', len(pdf) == n)
    graph=load(PUB/'verification/figure_evidence.json')
    check('Graphical abstract current numbers', np.isclose(graph['grain_delta_Mt_per_year'],gain) and
          np.isclose(graph['ET_delta_km3_per_year'],et) and
          np.isclose(graph['grain_delta_t_ha_per_year'],gain*1e6/means.mapped_rotation_area_ha.iloc[0]) and
          '+0.73' in (PUB/'figures/graphical_abstract.svg').read_text())
    book = pd.ExcelFile(PUB / 'tables/paper_data_and_tables.xlsx')
    check('Current data and table workbook', len(book.sheet_names) == 32 and
          all(name in book.sheet_names for name in ['Regional_annual_scenarios','Spatial_distribution','Grain_gain_reconciliation']))
    book_regions = pd.read_excel(book, 'Regional_annual_scenarios')
    check('Workbook current regional observations', len(book_regions) == len(yearly) and
          np.allclose(book_regions.grain_production_t, yearly.grain_production_t))
    result = {'verified_utc': datetime.now(timezone.utc).isoformat(), 'all_checks_passed': True,
              'check_count': len(checks), 'checks': checks, 'documents': doc_receipts,
              'selected_model': frozen['selected_version'], 'publication_source': 'current calibration and recalculated regional scenarios',
              'field_accuracy_criteria_pass': False, 'regional_results_classification': 'conditional research scenarios',
              'source_sha256': sha(Path(__file__))}
    (ROOT / 'verification/current_product_checks.json').write_text(json.dumps(result, indent=2) + '\n')
    print(f'{len(checks)} scientific-source and publication checks passed.')


if __name__ == '__main__':
    main()
