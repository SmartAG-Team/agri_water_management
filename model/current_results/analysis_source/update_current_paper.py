"""Bind the paper to the single selected crop fit and current regional results."""
from pathlib import Path
from copy import deepcopy
import hashlib
import json
import re
import shutil
import sys

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
P = ROOT / 'publication'
R = ROOT / 'regional'
sys.path.insert(0, str(P / 'analysis_source'))
from word_documents import manuscript, cover_signoff
from export_cited_references import main as export_references
from current_crop_narrative import apply as crop_narrative
from revise_manuscript_text import apply as editorial_revision, highlights as current_highlights


def load(path):
    return json.loads(Path(path).read_text())


def write(path, value):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(value, indent=2, ensure_ascii=False) + '\n')


def section(document, prefix):
    return next(b for b in document['blocks'] if b.get('heading', '').startswith(prefix))


def figure(document, label):
    return next(b for b in document['blocks'] if b.get('caption', '').startswith('Figure '+label+'.'))


def table(document, label):
    return next(b for b in document['blocks'] if b.get('table_caption', '').startswith('Table '+label+'.'))


def main():
    baseline = P / 'source_snapshots/previous_paper'
    baseline.mkdir(parents=True, exist_ok=True)
    for name in ['manuscript_blocks.json', 'supplementary_blocks.json', 'combined_package_blocks.json']:
        if not (baseline / name).exists():
            shutil.copy2(P / 'analysis_source' / name, baseline / name)
    article, supplement = crop_narrative(load(baseline/'manuscript_blocks.json'), load(baseline/'supplementary_blocks.json'))
    benchmark=pd.read_csv(P/'tables/main_benchmark_metrics.csv')
    benchmark['Partition']=benchmark.Partition.replace({'Calibration (2016–2018)':'2016–2018','Retrospective testing (2019)':'2019'})
    benchmark['Target']=benchmark.Target.replace({'Annual grain':'Grain','Aboveground biomass':'Biomass'})
    benchmark['Unit']=benchmark.Unit.replace({'t ha⁻¹ at 13% moisture':'t ha⁻¹','t ha⁻¹ dry mass':'t ha⁻¹'})
    benchmark.to_csv(P/'tables/main_benchmark_metrics.csv',index=False)
    table(article,'3').update(column_weights=[.5,.75,.8,.2,.5,.55,.55,.5,.45,.45],
        header_labels={'Partition':'Period','nRMSE_pct':'nRMSE (%)'})
    figure(article,'5')['caption']='Figure 5. Wuqiao calibration and retrospective testing. (a,b) Seasonal ET; (c,d) annual treatment-mean grain at 13% moisture. Circles show 2016–2018 treatments used to fit water coefficients; squares show 2019, withheld from the current fit and selection. Dashed lines denote 1:1 agreement. Horizontal grain bars are published SE from three field replicates, digitized with annual means from source Supplementary Figure S2. Simulated dry grain is converted once by division by 0.87. All 16 treatment seasons per crop are retained. Observed ET is the published 0–2 m water-balance estimate assuming negligible runoff and drainage. Maize treatment labels identify preceding wheat management; maize irrigation is identical within each year.'
    annual = pd.read_csv(R/'tables/regional_policy_annual_results.csv')
    later = annual[annual.harvest_year.between(2014, 2025)]
    means = later.groupby('policy').mean(numeric_only=True)
    c, u, t = [means.loc[name] for name in ['conventional', 'uniform_50pct', 'targeted_50pct']]
    gain = (t.grain_production_t-u.grain_production_t)/1e6
    etgain = (t.modeled_total_et_volume_m3-u.modeled_total_et_volume_m3)/1e9
    drainage = (t.bottom_drainage_volume_m3-u.bottom_drainage_volume_m3)/1e9
    paired = later.pivot(index='harvest_year', columns='policy', values='grain_production_t')
    loss_years = int((paired.targeted_50pct-paired.uniform_50pct < 0).sum())
    comparison = pd.read_csv(R/'tables/policy_comparison.csv')
    common = comparison[comparison.scope.eq('common_GRACE_available')].set_index('policy')
    q95, q98 = common.loc['adaptive_95'], common.loc['adaptive_98']
    bands = pd.read_csv(R/'tables/current_spatial_latitude_bands.csv')
    south, centre, north = [bands.iloc[i] for i in range(3)]
    loss_area = 100*bands.loc[bands.source_screened_delta_yield_kg_ha.lt(0), 'area_fraction'].sum()
    quota = pd.read_csv(R/'tables/training_selected_quotas.csv')
    display = quota[quota.grain_retention_target_pct.eq(95)].copy()
    display['Class'] = display.relative_class.str.replace('_', ' ').str.capitalize()
    display['Wheat irrigation (mm)'] = 300*display.selected_quota_fraction
    display['Maize irrigation (mm)'] = 80*display.selected_quota_fraction
    display = display.rename(columns={'selected_quota_fraction':'Irrigation fraction', 'training_grain_retention_pct':'Selection grain retention (%)'})
    display[['Class','Irrigation fraction','Wheat irrigation (mm)','Maize irrigation (mm)','Selection grain retention (%)']].to_csv(P/'tables/main_class_irrigation_candidates.csv',index=False)

    regional = means.reset_index()[['policy','grain_production_t','field_irrigation_m3','modeled_total_et_volume_m3','bottom_drainage_volume_m3','runoff_volume_m3']].copy()
    regional['Grain (Mt yr⁻¹)'] = regional.grain_production_t/1e6
    for source, label in [('field_irrigation_m3','Irrigation'),('modeled_total_et_volume_m3','ET'),('bottom_drainage_volume_m3','Drainage'),('runoff_volume_m3','Runoff')]:
        regional[label+' (km³ yr⁻¹)'] = regional[source]/1e9
    regional['Policy'] = regional.policy.str.replace('_',' ').str.capitalize()
    order=['conventional','uniform_25pct','targeted_25pct','uniform_50pct','targeted_50pct','uniform_75pct','targeted_75pct']
    regional=regional.set_index('policy').loc[order].reset_index()
    regional[['Policy','Grain (Mt yr⁻¹)','Irrigation (km³ yr⁻¹)','ET (km³ yr⁻¹)','Drainage (km³ yr⁻¹)','Runoff (km³ yr⁻¹)']].to_csv(P/'tables/main_policy_water_balance.csv',index=False)
    spatial = pd.read_csv(R/'tables/full_quota_spatial_metrics.csv')
    spatial = spatial[spatial.variable.isin(['yield_kg_ha','et_mm'])].copy()
    mass = spatial.variable.eq('yield_kg_ha')
    spatial.loc[mass,['rmse','bias','mean_actual_cell_prediction']] /= 1000
    spatial['Unit'] = spatial.variable.map({'yield_kg_ha':'t ha⁻¹','et_mm':'mm'})
    spatial['variable'] = spatial.variable.map({'yield_kg_ha':'Grain yield','et_mm':'ET'})
    spatial['crop'] = spatial.crop.str.capitalize()
    spatial.to_csv(P/'tables/supplement_spatial_metrics.csv',index=False)
    labels={'adaptive_95':'Storage–rainfall 95%','adaptive_98':'Storage–rainfall 98%',
            'rain_95_matched':'Rainfall 95%, matched','rain_98_matched':'Rainfall 98%, matched',
            'rain_95_available':'Rainfall 95%, available','rain_98_available':'Rainfall 98%, available',
            'conventional':'Conventional'}
    control=comparison.copy()
    control['Strategy']=control.policy.map(labels)
    control['Period']=control.scope.map({'common_GRACE_available':'Nine GRACE years','all_testing_years':'All twelve years'})
    control=control.rename(columns={'n_years':'Years','grain_retention_pct':'Grain retained (%)',
        'irrigation_reduction_mm':'Irrigation cut (mm)','et_reduction_mm':'ET cut (mm)',
        'annual_minimum_grain_retention_pct':'Minimum grain (%)','field_irrigation_mm':'Irrigation (mm)',
        'actual_ET_mm':'ET (mm)'})
    control[control.policy.isin(['adaptive_95','adaptive_98','conventional'])][['Strategy','Period','Years','Grain retained (%)','Irrigation cut (mm)','ET cut (mm)','Minimum grain (%)']].to_csv(P/'tables/supplement_continuous_adaptive.csv',index=False)
    control[control.policy.ne('conventional')][['Strategy','Period','Grain retained (%)','Irrigation (mm)','ET (mm)']].to_csv(P/'tables/supplement_grace_rainfall_comparison.csv',index=False)
    latitude = pd.DataFrame({'Latitude band':bands.latitude_band,'Area share (%)':100*bands.area_fraction,
        'Targeted irrigation fraction':bands.source_screened_targeted_mean_quota,
        'Grain change (t ha⁻¹)':bands.source_screened_delta_yield_kg_ha/1000,
        'ET change (mm)':bands.source_screened_delta_modeled_total_et_mm,
        'Grain contribution (Mt)':bands.source_screened_grain_gain_t/1e6})
    latitude.to_csv(P/'tables/supplement_spatial_latitude_bands.csv',index=False)
    for origin, label, dest in [('Figure_3_policy_tradeoffs','6','Figure_3_policy_tradeoffs'),
                               ('Figure_7_current_spatial_policy_outcomes','7','Figure_7_current_spatial_policy_outcomes'),
                               ('adaptive_irrigation_testing','8','Figure_8_continuous_class_irrigation')]:
        if not (P/'verification/visual_revision_20261008.json').exists():
            for ext in ['png','pdf']:
                shutil.copy2(R/f'figures/{origin}.{ext}',P/f'figures/closed_axes/{dest}.{ext}')
        figure(article,label)['figure']=f'figures/closed_axes/{dest}.png'
    for origin, label, dest in [('all_quota_spatial_errors','S1','all_quota_spatial_errors'),
                               ('continuous_soil_storage','S3','Figure_S14_continuous_soil_storage')]:
        if not (P/'verification/panel_title_removal_20261008.json').exists():
            for ext in ['png','pdf']:
                shutil.copy2(R/f'figures/{origin}.{ext}',P/f'figures/closed_axes/{dest}.{ext}')
        figure(supplement,label)['figure']=f'figures/closed_axes/{dest}.png'
    if not (P/'verification/panel_title_removal_20261008.json').exists():
        for ext in ['png','pdf']:
            shutil.copy2(ROOT/f'calibration/figures/Field_grain_yield.{ext}',P/f'figures/current_annual_field_grain.{ext}')
    figure(supplement,'S2')['figure']='figures/current_annual_field_grain.png'

    section(article,'2.1.')['paragraphs']=[
        'Inherited multisite growth estimates and documented irrigation treatments constrain complementary crop responses (Figure 1). The Wuqiao 2016–2018 treatments estimate eight water coefficients per crop; 2019 and the original station partitions supply retrospective and conditional comparisons. Daily weather, mapped rotation area and soil profiles define 32 regional representatives.',
        'One frozen current parameter set drives continuous crops and fallows under five irrigation levels. Uniform and optimized allocations share regional irrigation budgets; storage–rainfall and rainfall-only strategies share selection periods and matched monitoring availability. Regional comparisons quantify grain production, field irrigation, actual ET and profile drainage during 2014–2025.']
    figure(article,'1')['caption']='Figure 1. Overall study workflow. Multisite growth estimates are retained while documented Wuqiao 2016–2018 treatments constrain water parameters; 2019 provides retrospective testing. The same frozen parameters drive continuous regional rotations, spatial allocation and rainfall–storage strategies. Regional outcomes during 2014–2025 are conditional retrospective scenarios.'
    intro = section(article,'1. Introduction')
    intro['paragraphs'][-1]='The objective was to quantify how spatial allocation and rainfall–storage conditioning influence grain production and water use in the North China Plain wheat–maize rotation. Open Crop Model combines inherited multisite growth estimates with water coefficients calibrated against documented Wuqiao irrigation treatments. Retrospective field and conditional station comparisons characterize the resulting crop responses. Continuous regional simulations compare uniform and optimized allocations under identical irrigation constraints, and storage–rainfall strategies with rainfall-only controls selected from historical responses. The analysis assesses production retention, consumptive use, profile drainage and the distribution of production changes across locations and years.'
    section(article,'2.7.')['paragraphs'] = [section(article,'2.7.')['paragraphs'][0],
        'The same current process equations and crop parameters generate conventional, uniform and targeted responses. Scenario selection uses 1997–2013 weather responses, while outcomes are compared during 2014–2025. Current crop fitting includes 2016–2018 field data; this calendar overlap makes the regional assessment a retrospective scenario comparison rather than time-forward crop validation.',
        'Spatial contrasts at the common 50% irrigation cut use frozen shares of complete continuous irrigation histories. Their area-weighted grain, ET and drainage differences map from 32 representatives to 3,641 source cells. Fixed southern (32–36°N), central (36–38°N) and northern (38–41°N) bands summarize the distribution. Counts of positive annual contrasts describe consistency within the modeled record (Figure 7; Table S9).']
    section(article,'3.3.')['paragraphs']=[
        f'Conventional irrigation supplied {c.field_irrigation_m3/1e9:.3f} km³ yr⁻¹ and produced {c.grain_production_t/1e6:.3f} million t yr⁻¹ of simulated dry grain during 2014–2025. At a 50% irrigation cut, uniform and targeted allocation each supplied {u.field_irrigation_m3/1e9:.3f} km³ yr⁻¹. Targeting retained {t.grain_production_t/1e6:.3f} million t yr⁻¹, compared with {u.grain_production_t/1e6:.3f} under uniform cuts, a difference of {gain:.3f} million t yr⁻¹ ({100*gain/(u.grain_production_t/1e6):.2f}% of uniform production; Table 4).',
        f'Targeted and uniform crop-plus-fallow ET were {t.modeled_total_et_volume_m3/1e9:.3f} and {u.modeled_total_et_volume_m3/1e9:.3f} km³ yr⁻¹, respectively, compared with {c.modeled_total_et_volume_m3/1e9:.3f} under conventional irrigation. Targeting increased ET by {etgain:.3f} km³ yr⁻¹ under the same irrigation constraint. Profile drainage was {t.bottom_drainage_volume_m3/1e9:.3f} versus {u.bottom_drainage_volume_m3/1e9:.3f} km³ yr⁻¹, and runoff was {t.runoff_volume_m3/1e9:.3f} versus {u.runoff_volume_m3/1e9:.3f} km³ yr⁻¹ (Table 4; Figure 6).',
        f'The targeted-minus-uniform grain contrast was positive in {12-loss_years} of 12 years. Spatial gains were concentrated in the central and northern latitude bands, contributing {centre.source_screened_grain_gain_t/1e6:.3f} and {north.source_screened_grain_gain_t/1e6:.3f} million t yr⁻¹. The southern band contributed {south.source_screened_grain_gain_t/1e6:.3f} million t yr⁻¹ and covered {loss_area:.1f}% of mapped rotation area (Figure 7; Table S9).']
    table(article,'4').update(table='tables/main_policy_water_balance.csv',columns=['Policy','Grain (Mt yr⁻¹)','Irrigation (km³ yr⁻¹)','ET (km³ yr⁻¹)','Drainage (km³ yr⁻¹)','Runoff (km³ yr⁻¹)'],
        table_note='Current-parameter conditional means during 2014–2025. Grain is dry mass and ET includes crops and fallows. Uniform and targeted policies deliver identical irrigation at each reduction. Drainage leaves the modeled 2-m profile and does not estimate aquifer recharge.')
    figure(article,'6')['caption']='Figure 6. Current-parameter uniform and targeted irrigation scenarios during 2014–2025. (a,b) Annual grain production and crop-plus-fallow actual ET; error bars are between-year standard deviations, not confidence intervals. (c) Paired annual targeted grain advantage at identical irrigation volumes. (d) Grain–ET combinations with large symbols for means and small symbols for individual years. Allocations are selected using 1997–2013 responses and remain fixed.'
    figure(article,'7')['caption']='Figure 7. Current spatial consequences of targeting at a common 50% irrigation cut. Panels show (a) targeted irrigation fractions, (b) uniform fractions, (c) targeted-minus-uniform dry grain, (d) actual ET, (e) the number of positive grain-contrast years and (f) latitude-band grain contrasts. Source cells inherit their representative responses without interpolation; fractions are area shares of complete simulated histories. Means cover 2014–2025.'
    section(article,'3.4.')['paragraphs']=[
        'Four antecedent storage–rainfall classes occurred during 2003–2013. Nine later years had eligible GRACE observations, all in the lower-storage category; three missing-storage years received conventional irrigation. At the 95% grain target, both storage strata selected fraction 0.75 under lower rainfall and 0.50 under higher rainfall. The 98% target selected 1.00 and 0.75, respectively (Table 5).',
        f'Continuous strategies selected at the 95% and 98% targets retained {q95.grain_retention_pct:.2f}% and {q98.grain_retention_pct:.2f}% of conventional grain over the nine class-available years. Field-irrigation reductions were {q95.irrigation_reduction_mm:.2f} and {q98.irrigation_reduction_mm:.2f} mm per rotation; corresponding ET reductions were {q95.et_reduction_mm:.2f} and {q98.et_reduction_mm:.2f} mm (Figure 8; Table S7). The 95% strategy retained at least 95% in {int(q95.years_retaining_at_least_95pct)} of nine years, with a minimum of {q95.annual_minimum_grain_retention_pct:.2f}%.',
        'Storage–rainfall and matched rainfall-only strategies selected identical irrigation fractions and produced identical grain, ET, drainage and irrigation histories at both targets. Applying rainfall-only decisions during missing-GRACE years changed preceding management and subsequent soil-water states; those histories are reported separately (Table S8).']
    table(article,'5').update(table='tables/main_class_irrigation_candidates.csv',columns=list(display[['Class','Irrigation fraction','Wheat irrigation (mm)','Maize irrigation (mm)','Selection grain retention (%)']].columns),
        table_note='Fractions meet the 95% class-mean grain target during selection years 2003–2013. Current shared crop parameters apply throughout. Both storage strata select identical fractions within each rainfall group; higher-storage classes have no later-period exposure.')
    figure(article,'8')['caption']='Figure 8. Continuous current-parameter irrigation strategies during 2014–2025. Panels show field irrigation, grain retained, crop-plus-fallow ET reduction and annual grain–irrigation trade-offs. Circles and squares distinguish the 95% and 98% selection targets. Missing-storage years use conventional fallback. Soil water carries continuously through crops, fallows and annual quota changes; the strategies are retrospective conditional scenarios.'
    section(article,'4.1.')['paragraphs']=[
        'Spatial allocation rewards differences in marginal grain response to irrigation across climate and soil conditions. Field schedules show diminishing production gains as water supply increases [CITE:yang2022_long_term_irrigation], while regional optimization studies identify scope for redistributing deliveries [CITE:jiang2016_regional_optimization|wu2019_spatial_irrigation]. Here, allocations use complete continuous histories rather than interpolated single-season response curves. Their comparison holds total applied irrigation constant and therefore isolates the consequences of where that delivery budget is assigned.',
        f'At the common 50% cut, targeting retains an additional {gain:.3f} Mt yr⁻¹, with a positive contrast in all twelve comparison years. Regional gain nevertheless coexists with a negative mean contrast across the southern latitude band, which accounts for {loss_area:.1f}% of mapped area. The objective protects area-weighted mean production, while permitting production losses at individual locations. A regional production target therefore requires local safeguards if the farms receiving larger reductions are to share the benefits of redistribution.',
        'Spatial clustering also distinguishes modeled allocation gains from crop-model representation error. Eight purposive actual-cell simulations quantify discrepancies under the same histories; they do not sample all environmental or parameter uncertainty. Their role is to assess the representative-unit approximation, while the field test addresses a different source of uncertainty: the crop response functions being optimized.']
    section(article,'4. Discussion')['paragraphs']=[
        f'The strategies depend on the modeled response of rotation grain production to water supply. Irrigation changes profile storage, rooted-layer uptake, crop stress and grain formation; carried water makes the response depend on preceding management. Spatial targeting redistributes a fixed delivery budget, while rainfall–storage strategies select an annual quota. At the same 50% cut, the current parameter set retains {gain:.3f} Mt yr⁻¹ more grain and {etgain:.3f} km³ yr⁻¹ more ET than uniform allocation (Table 4; Figure 6). Their fixed application dates and conditional crop responses distinguish these comparisons from operational soil-moisture-triggered scheduling. Numerical closure verifies conservation; field response and soil-state evidence govern transfer.']
    section(article,'4.3.')['paragraphs']=[
        f'Targeting at the common 50% cut changes ET by {etgain:+.3f} km³ yr⁻¹ and drainage by {drainage:+.3f} km³ yr⁻¹ relative to uniform reduction. Under identical rainfall and irrigation, changes in consumption are balanced by drainage, runoff and stored water. Maintaining more grain can consequently increase consumption even when the delivered irrigation budget is unchanged. Irrigation productivity and consumptive-water savings are separate objectives [CITE:perry2009_water_accounting|grafton2018_irrigation_paradox].',
        'The continuous profile retains the effects of preceding crops, rainfall and delivery decisions. Drainage leaving 2 m enters a deeper vadose zone whose storage and transit are outside the model [CITE:wu2023_vadose_zone]. Neither reduced applied irrigation nor reduced profile drainage identifies an equivalent change in aquifer recharge or pumping. Water-accounting boundaries determine which saving a management comparison measures.',
        'Regional GRACE anomalies aggregate groundwater with other terrestrial stores. Their effective observation scale exceeds the displayed grid scale [CITE:save2016_csr_mascons|scanlon2016_grace_mascons]. Recent observed aquifer recovery reflects broader interventions and hydroclimatic conditions [CITE:long2025_aquifer_recovery|xu2026_groundwater_recovery]. Crop-model allocation scenarios can inform production and consumption trade-offs without attributing those regional storage changes to individual irrigation strategies.']
    section(article,'4.4.')['paragraphs']=[
        'A hydrological indicator adds decision value when it changes the action selected under the relevant production and water constraints. GRACE identifies regional storage conditions, while antecedent rainfall describes more local supply; their different scales motivate a joint classification [CITE:save2016_csr_mascons|scanlon2016_grace_mascons]. This physical distinction alone does not establish an improvement in the available irrigation choices.',
        'The current storage–rainfall and rainfall-only strategies select identical fractions within corresponding rainfall groups at both grain targets. Their complete matched-coverage histories therefore have identical crop states and water fluxes. In this five-level quota design, the storage split characterizes conditions without altering the preferred action. All nine common comparison years occupy the lower-storage category, leaving the higher-storage selections without later-period exposure.',
        'Applying rainfall-only decisions during GRACE gaps tests monitoring coverage as well as the grouping criterion. Earlier irrigation differences alter the storage inherited by later crops, so common-year production can differ even when mean irrigation in those years is the same. Matched coverage supplies the valid incremental-information comparison. Finer quota choices, different thresholds or prospective availability would constitute different management experiments and require their own selection and evaluation.']
    reliability = section(article,'4.5.')
    reliability['paragraphs'] = reliability['paragraphs'][:2]+[
        f'The {q95.grain_retention_pct:.2f}% and {q98.grain_retention_pct:.2f}% period-mean retention values conceal annual minima of {q95.annual_minimum_grain_retention_pct:.2f}% and {q98.annual_minimum_grain_retention_pct:.2f}% (Figure 8). A mean production target does not specify a shortfall probability or protect every farm. Operational constraints can include local production floors, acceptable shortfall frequency and a reserve entering the next crop. Forecast-based irrigation work illustrates explicit risk constraints on water stress [CITE:roy2021_irrigation_risk]. Regional allocation supplies a delivery budget; field scheduling additionally depends on current soil water, forecasts and application timing.']
    section(article,'Abstract')['paragraphs']=[
        'Irrigation sustains wheat–maize production in the North China Plain, where rainfall and soil-water carry-over create uneven irrigation returns. A process-based crop model combined inherited multisite growth estimates with water coefficients calibrated against documented 2016–2018 Wuqiao treatments and retrospectively tested in 2019. Continuous regional rotations compared uniform and targeted allocation under identical irrigation budgets and storage–rainfall strategies against matched rainfall-only controls. Allocations and strategies were selected from historical simulated responses and compared during 2014–2025. '
        f'At a 50% reduction in regional field irrigation, targeting retained {gain:.2f} million t yr⁻¹ more simulated dry grain than uniform cuts, with {etgain:.2f} km³ yr⁻¹ greater evapotranspiration. Central and northern production gains outweighed losses in the southern latitude band. Over nine GRACE-available years, strategies selected at 95% and 98% grain targets retained {q95.grain_retention_pct:.2f}% and {q98.grain_retention_pct:.2f}% of conventional production. The 95% strategy reduced irrigation by {q95.irrigation_reduction_mm:.2f} mm and ET by {q95.et_reduction_mm:.2f} mm per rotation. Matched storage–rainfall and rainfall-only strategies selected identical quotas and outcomes. These conditional comparisons show that production-oriented targeting can sustain greater consumption under a shared delivery budget. Regional irrigation management consequently requires explicit consumptive-use constraints, local production safeguards and evaluation of whether additional hydrological information changes decisions.']
    section(article,'5. Conclusions')['paragraphs']=[
        f'Under an identical 50% reduction in field irrigation, current-parameter targeting retained {gain:.3f} Mt yr⁻¹ more simulated dry grain and {etgain:.3f} km³ yr⁻¹ more crop-plus-fallow ET than uniform reduction. Central and northern gains outweighed southern losses; a regional production objective therefore requires local safeguards and a separate consumption constraint.',
        f'Continuous strategies selected at 95% and 98% grain targets retained {q95.grain_retention_pct:.2f}% and {q98.grain_retention_pct:.2f}% during nine GRACE-available years. The 95% strategy reduced irrigation by {q95.irrigation_reduction_mm:.2f} mm and ET by {q95.et_reduction_mm:.2f} mm per rotation. Matched rainfall-only controls produced identical actions and outcomes, so the storage classification supplied no additional decision benefit in this configuration.',
        'Documented treatment management constrains the current water fit, while the single retrospective field test and conditional station observations delimit transfer. Regional results describe the selected crop model, fixed schedules and spatial support; they do not establish groundwater recovery or operational field performance.']
    section(article,'Data and code availability')['paragraphs']=['Model code, frozen calibration and regional inputs, observation subsets and quality exclusions, weather, calibrated parameters, predictions, evaluation metrics and figure-data tables are available at https://github.com/SmartAG-Team/agri_water_management. The repository includes a standalone Open Crop Model package, pinned dependencies, reproduction commands and SHA-256 provenance. Reproduction uses the included inputs without another model checkout or access to the historical workspace. Original datasets retain their provider licences and attribution requirements, documented in the repository source catalogue.']

    section(supplement,'S5.')['paragraphs']=[
        'All nine common GRACE-available years occupy lower-storage classes; three other years use conventional fallback. The selected current quotas are identical across storage strata within each rainfall group. Matched rainfall-only and storage–rainfall histories have identical actions, crop states and outcomes throughout the modeled record.',
        'Available-rainfall controls apply reduced quotas during monitoring gaps and therefore create different carried water states before later common years. These history effects are separate from the value of the storage classification. Table S7 reports all-year and common-year retention and water use; Table S8 reports the matched and available-rainfall controls.']
    table(supplement,'S7').update(table='tables/supplement_continuous_adaptive.csv',columns=['Strategy','Period','Years','Grain retained (%)','Irrigation cut (mm)','ET cut (mm)','Minimum grain (%)'],table_note='Current continuous histories. Grain retention is a ratio of period-mean regional dry-grain production to conventional irrigation. ET includes crops and fallows. All-year summaries include conventional fallback; common-year summaries use the nine GRACE-available years.')
    table(supplement,'S8').update(table='tables/supplement_grace_rainfall_comparison.csv',columns=['Strategy','Period','Grain retained (%)','Irrigation (mm)','ET (mm)'],table_note='Matched controls share the complete GRACE-availability and fallback calendar. Available-rainfall controls act during missing-storage years and alter subsequent soil states. All values use current parameters; differences do not identify causal groundwater responses.')
    section(supplement,'S6.')['paragraphs']=['Original Yucheng harvest sources contain identical six-replicate maize vectors in 2012 and 2017 under different date and cultivar labels. Both groups are retained as raw evidence and excluded from scored grain and harvest biomass. Eligibility depends on source authenticity rather than model residuals. Current plots preserve observed quadrat variation and the original whole site-year partitions (Figure S4).']
    section(supplement,'S8.')['paragraphs']=['Frozen area shares weight complete current irrigation histories before spatial aggregation. Dry-grain differences in kg ha⁻¹ multiply represented hectares and divide by 1,000 for tonnes; water depths multiply hectares by 10 for cubic metres. Independent sums reproduce regional contrasts. Latitude bands describe the distribution of gains and losses, while the 3,641 mapped cells inherit 32 simulated representative responses (Table S9).']
    table(supplement,'S12').update(table='tables/supplement_spatial_latitude_bands.csv',columns=list(latitude.columns),table_note='Current targeted-minus-uniform means at the same 50% regional irrigation cut. Latitude bands are reporting partitions. Source cells inherit representative responses and are not independent observations.')
    supplement['blocks']=[b for b in supplement['blocks'] if not b.get('heading','').startswith('S7.') and not b.get('table_caption','').startswith(('Table S9.','Table S10.','Table S11.')) and not b.get('caption','').startswith('Figure S5.')]
    for b in supplement['blocks']:
        if b.get('heading','').startswith('S8.'):b['heading']=b['heading'].replace('S8.','S7.',1)
        elif b.get('heading','').startswith('S9.'):b['heading']=b['heading'].replace('S9.','S8.',1)
    for document in [article,supplement]:
        for b in document['blocks']:
            for key in ['caption','table_caption','table_note','text']:
                if key in b:b[key]=re.sub(r'\bTable S12\b','Table S9',b[key])
            b['paragraphs']=[re.sub(r'\bTable S12\b','Table S9',p) for p in b.get('paragraphs',[])]
            def renumber(text):
                return re.sub(r'\bFigure S(\d+)\b',lambda m:'Figure S'+str(int(m.group(1))-1) if int(m.group(1))>=6 else m.group(0),text)
            for key in ['caption','table_note','text']:
                if key in b:b[key]=renumber(b[key])
            b['paragraphs']=[renumber(p).replace('Figures S7–S12','Figures S6–S11') for p in b.get('paragraphs',[])]
    table(supplement,'S4')['table_note']='Current model aggregation discrepancies against eight actual-cell simulations under all five irrigation fractions. Grain is converted from kg ha⁻¹ to t ha⁻¹ in every displayed grain-valued column. Exactly zero irrigation discrepancies are omitted from display; full inputs and metrics remain in the current-results product.'
    figure(supplement,'S1')['caption']='Figure S1. Current crop-model spatial aggregation errors under all five irrigation fractions in eight purposively selected cells during 2014–2025. Actual-cell and representative predictions share the current native model and crop parameters. These discrepancies assess spatial representation rather than field-observation accuracy.'
    figure(supplement,'S3')['caption']='Figure S3. Current simulated 0–2 m soil-water storage in two representative units. State carries continuously through wheat, maize, fallows and annual irrigation choices. Displayed histories compare conventional irrigation with selected current strategies; they are model states rather than observed groundwater levels.'
    figure(article,'2')['caption']=figure(article,'2')['caption'].replace('The independent Wuqiao irrigation experiment', 'The documented Wuqiao irrigation experiment')
    table(supplement,'S3')['column_weights']=[2.6,1.55,.9,.9]
    for document in [article,supplement]:
        for block in document['blocks']:
            if 'caption' in block:
                block['caption']=re.sub(r'\blai\b','LAI',block['caption'])
                block['caption']=re.sub(r'\bet\b','ET',block['caption'])

    article,supplement=editorial_revision(article,supplement)
    cover = [
        'Dear Editor,',
        'Please consider the research article “Irrigation strategies under contrasting water-storage and rainfall conditions in the North China Plain” for publication in Agricultural Water Management.',
        'Reducing irrigation while sustaining grain production requires balancing water consumption and the distribution of production losses. Using a process-based crop model, the study compares uniform and spatially targeted irrigation under equal water budgets and tests whether satellite water-storage information improves rainfall-based irrigation decisions.',
        f'At the same 50% irrigation reduction, targeting retained {100*(t.grain_production_t/u.grain_production_t-1):.2f}% more simulated grain than uniform cuts, but increased evapotranspiration and reduced grain production in some locations. Adding regional storage information did not change irrigation quotas or outcomes relative to matched rainfall-only strategies within the tested design.',
        'These model-based findings support irrigation strategies that combine production goals with limits on consumptive use and safeguards for local production. Their focus on irrigation allocation and crop water use makes the study well suited to Agricultural Water Management.',
        'Thank you for considering this manuscript.',
        'Sincerely,',
        *cover_signoff(),
    ]
    write(P/'analysis_source/cover_letter_paragraphs.json',cover)
    from docx import Document
    from docx.shared import Pt
    document=Document(P/'documents/cover_letter.docx')
    for paragraph in list(document.paragraphs):paragraph._p.getparent().remove(paragraph._p)
    for text in cover:
        q=document.add_paragraph(text);q.paragraph_format.line_spacing=1.1;q.paragraph_format.space_after=Pt(8)
    document.save(P/'documents/cover_letter.docx')
    highlights=current_highlights()
    assert len(highlights)==5 and all(len(item)<=85 for item in highlights)
    write(P/'analysis_source/highlights_paragraphs.json',highlights)
    document=Document(P/'documents/highlights.docx')
    for q in list(document.paragraphs):q._p.getparent().remove(q._p)
    for text in highlights:document.add_paragraph(text,style='List Bullet')
    document.save(P/'documents/highlights.docx')
    write(P/'analysis_source/manuscript_blocks.json',article)
    write(P/'analysis_source/supplementary_blocks.json',supplement)
    combined=load(baseline/'combined_package_blocks.json')
    marker=next(i for i,b in enumerate(combined['blocks']) if b.get('heading')=='Manuscript')
    prefix=deepcopy(combined['blocks'][:marker+1])
    next(b for b in prefix if b.get('heading')=='Cover letter')['paragraphs']=cover
    next(b for b in prefix if b.get('heading')=='Highlights')['paragraphs']=highlights
    next(b for b in prefix if b.get('heading')=='Highlights')['bullet_list']=True
    combined['blocks']=prefix+article['blocks']+[{'heading':'Supplementary material','page_break':True}]+supplement['blocks']
    write(P/'analysis_source/combined_package_blocks.json',combined)
    build={}
    for name, source in [('manuscript','manuscript_blocks'),('supplementary_material','supplementary_blocks'),('manuscript_package','combined_package_blocks')]:
        build[name+'.docx']=manuscript(P/f'analysis_source/{source}.json',P/f'documents/{name}.docx')
    write(P/'verification/document_build_receipt.json',build)
    export_references()
    binding={'selected_model':'management_refit','calibration_source':'../calibration','regional_source':'../regional',
        'abstract_words':len(section(article,'Abstract')['paragraphs'][0].split()),'cover_words':len(' '.join(cover).split()),
        'grain_advantage_Mt':gain,'ET_advantage_km3':etgain,'grain_retention95':float(q95.grain_retention_pct),
        'grain_retention98':float(q98.grain_retention_pct),'main_figures':sum('figure'in b for b in article['blocks']),
        'supplement_figures':sum('figure'in b for b in supplement['blocks']),
        'main_tables':sum('table'in b for b in article['blocks']),'supplement_tables':sum('table'in b for b in supplement['blocks']),
        'grain_advantage_t_ha_yr':gain*1e6/float(c.mapped_rotation_area_ha),
        'ET_advantage_mm_yr':etgain*1e9/(float(c.mapped_rotation_area_ha)*10),
        'highlights':5,'revision_date':'2026-10-08','cartographic_boundary_sources':'publication/source_snapshots/cartography',
        'source_sha256':{str(f.relative_to(ROOT)):hashlib.sha256(f.read_bytes()).hexdigest() for f in [ROOT/'calibration/parameters/frozen_model.json',R/'tables/regional_policy_annual_results.csv',R/'tables/policy_comparison.csv',Path(__file__)]}}
    write(P/'verification/current_publication_binding.json',binding)
    print(json.dumps(binding,indent=2))


if __name__=='__main__':main()
